"""PowerPoint (PPTX) parser using python-pptx with exact slide numbering and figure extraction."""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.tutor.models import SlideSource
from app.kb.models import FigureUnit
from app.kb.media_store import MediaStore

logger = logging.getLogger(__name__)


@dataclass
class ExtractedSlide:
    slide_number: int
    title: str
    body_text: str
    speaker_notes: str
    full_text: str
    figures: list[FigureUnit] = field(default_factory=list)


@dataclass
class PptParseResult:
    file_name: str
    course_id: str
    total_slides: int
    slides: list[ExtractedSlide]
    figures: list[FigureUnit]


class PptParser:
    """Extracts slide text, speaker notes, and embedded diagrams from PPTX presentations."""

    def __init__(
        self,
        media_store: MediaStore | None = None,
        min_image_dimension: int = 80,
    ) -> None:
        self.media_store = media_store
        self.min_image_dimension = min_image_dimension

    def parse(self, file_path: Path | str, course_id: str) -> PptParseResult:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"PPTX file not found: {path}")

        file_name = path.name
        prs = Presentation(str(path))
        extracted_slides: list[ExtractedSlide] = []
        all_figures: list[FigureUnit] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1  # 1-indexed

            # 1. Slide Title
            title_text = ""
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title_text = slide.shapes.title.text_frame.text.strip()

            # 2. Extract shape text (text boxes, shapes, tables)
            text_parts: list[str] = []
            slide_figures: list[FigureUnit] = []
            img_counter = 0

            for shape in slide.shapes:
                # Text frames
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t and t != title_text:
                        text_parts.append(t)

                # Tables
                elif shape.has_table:
                    table_rows: list[str] = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        table_rows.append(" | ".join(row_cells))
                    if table_rows:
                        text_parts.append("\n".join(table_rows))

                # Embedded Pictures
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        img_counter += 1
                        image_bytes = shape.image.blob
                        pil_img = Image.open(io.BytesIO(image_bytes))
                        width, height = pil_img.size

                        if (
                            width >= self.min_image_dimension
                            and height >= self.min_image_dimension
                        ):
                            fig_id = f"fig_{path.stem}_s{slide_num}_{img_counter}"
                            rel_path = ""
                            if self.media_store:
                                rel_path = self.media_store.save_image(
                                    course_id=course_id,
                                    figure_id=fig_id,
                                    image=pil_img,
                                    image_format=pil_img.format or "PNG",
                                )

                            figure = FigureUnit(
                                id=fig_id,
                                course_id=course_id,
                                source=SlideSource(
                                    type="slide",
                                    file_name=file_name,
                                    slide=slide_num,
                                ),
                                image_rel_path=rel_path,
                                caption=title_text or f"Slide {slide_num} Diagram {img_counter}",
                                explanation="",
                                metadata={
                                    "width": width,
                                    "height": height,
                                    "format": shape.image.ext,
                                },
                            )
                            slide_figures.append(figure)
                            all_figures.append(figure)
                    except Exception as e:
                        logger.warning(
                            f"failed_extracting_slide_image_s{slide_num}: {e}"
                        )

            # 3. Speaker notes
            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()

            # 4. Synthesize complete slide text
            full_blocks: list[str] = []
            if title_text:
                full_blocks.append(f"## {title_text}")
            if text_parts:
                full_blocks.append("\n\n".join(text_parts))
            if notes_text:
                full_blocks.append(f"[Speaker Notes: {notes_text}]")

            full_text = "\n\n".join(full_blocks).strip()

            extracted_slides.append(
                ExtractedSlide(
                    slide_number=slide_num,
                    title=title_text,
                    body_text="\n\n".join(text_parts).strip(),
                    speaker_notes=notes_text,
                    full_text=full_text,
                    figures=slide_figures,
                )
            )

        if self.media_store and all_figures:
            self.media_store.add_figures(all_figures)

        logger.info(
            "pptx_parsed_successfully",
            extra={
                "file_name": file_name,
                "slides": len(extracted_slides),
                "figures": len(all_figures),
            },
        )

        return PptParseResult(
            file_name=file_name,
            course_id=course_id,
            total_slides=len(prs.slides),
            slides=extracted_slides,
            figures=all_figures,
        )
