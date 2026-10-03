"""End-to-end integration test verifying ingestion of PDF, PPTX, and Video,
followed by vector retrieval and prerequisite graph validation.
"""

import tempfile
from pathlib import Path

import fitz  # PyMuPDF
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.tutor.knowledge_base import KnowledgeBase
from app.tutor.vector_store import InMemoryVectorStore
from app.tutor.providers.embeddings import HashingEmbeddingProvider
from app.tutor.retriever import Retriever
from app.kb.media_store import MediaStore
from app.kb.graph_store import GraphStore
from app.ingestion.pipeline import IngestionPipeline


def test_end_to_end_multimodal_ingestion():
    with tempfile.TemporaryDirectory() as tmpdir:
        root = Path(tmpdir)
        media_dir = root / "figures"
        graph_dir = root / "graphs"
        data_dir = root / "kb_runtime"

        # 1. Setup Knowledge Base and stores
        vector_store = InMemoryVectorStore()
        embedder = HashingEmbeddingProvider()
        kb = KnowledgeBase(store=vector_store, embedder=embedder, persist_dir=data_dir)
        media_store = MediaStore(base_dir=media_dir)
        graph_store = GraphStore(persist_dir=graph_dir)

        pipeline = IngestionPipeline(
            kb=kb,
            media_store=media_store,
            graph_store=graph_store,
            api_key=None,
        )

        course_id = "test_multimodal_course"

        # 2. Create sample PDF with text and embedded image
        pdf_path = root / "textbook_ch1.pdf"
        doc = fitz.open()
        page = doc.new_page(width=595, height=842)
        page.insert_text(
            (50, 72),
            "Photosynthesis in Higher Plants\n"
            "Chloroplasts are the primary organelles of photosynthesis.\n"
            "The light reactions produce ATP and NADPH in the thylakoid membranes.",
            fontsize=12,
        )

        # Create sample diagram image
        img = Image.new("RGB", (200, 200), color=(73, 109, 137))
        d = ImageDraw.Draw(img)
        d.text((10, 10), "Thylakoid Diagram", fill=(255, 255, 0))
        img_path = root / "temp_diagram.png"
        img.save(img_path)

        # Insert image into page
        page.insert_image(fitz.Rect(50, 150, 250, 350), filename=str(img_path))
        doc.save(str(pdf_path))
        doc.close()

        # Ingest PDF
        pdf_summary = pipeline.ingest_file(course_id=course_id, file_path=pdf_path)
        assert pdf_summary.units_count >= 1
        assert pdf_summary.figures_count >= 1
        assert not pdf_summary.errors

        # 3. Create sample PPTX slide deck
        pptx_path = root / "lecture_slides.pptx"
        prs = Presentation()
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)

        # Add title textbox
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        tf = txBox.text_frame
        tf.text = "Calvin Cycle and Carbon Fixation"

        # Add notes
        notes_slide = slide.notes_slide
        notes_text_frame = notes_slide.notes_text_frame
        notes_text_frame.text = "Explain the role of RuBisCO in carbon fixation."

        prs.save(str(pptx_path))

        # Ingest PPTX
        pptx_summary = pipeline.ingest_file(course_id=course_id, file_path=pptx_path)
        assert pptx_summary.units_count >= 1
        assert not pptx_summary.errors

        # 4. Create sample Video transcript
        video_dummy = root / "lecture_video.mp4"
        video_dummy.write_bytes(b"dummy_mp4_bytes")
        vtt_path = root / "lecture_video.vtt"
        vtt_path.write_text(
            """WEBVTT

00:03:00.000 --> 00:04:30.000
Chemiosmosis links the electron transport chain to the phosphorylation of ADP.
Protons pump across the membrane creating an electrochemical gradient.
""",
            encoding="utf-8",
        )

        # Ingest Video with companion transcript
        video_summary = pipeline.ingest_file(
            course_id=course_id,
            file_path=video_dummy,
            subtitle_path=vtt_path,
        )
        assert video_summary.units_count >= 1
        assert not video_summary.errors

        # 5. Verify Knowledge Base retrieval using Member 2's Retriever
        retriever = Retriever(store=kb.store, embedder=kb.embedder, default_top_k=5)
        hits = retriever.retrieve(course_id=course_id, query="thylakoid ATP NADPH")
        assert len(hits) >= 1
        top_hit = hits[0]
        assert top_hit.course_id == course_id
        assert top_hit.source.type == "pdf"
        assert top_hit.source.page == 1

        # Search for video segment
        video_hits = retriever.retrieve(course_id=course_id, query="Chemiosmosis gradient")
        assert len(video_hits) >= 1
        vid_hit = next((h for h in video_hits if h.source.type == "video"), None)
        assert vid_hit is not None
        assert vid_hit.source.timestamp_start == "00:03:00"

        # 6. Verify Prerequisite Graph & Learning Path (Member 4 Contract)
        graph = graph_store.get_graph(course_id)
        assert graph is not None
        assert len(graph.nodes) >= 1
        order = graph_store.get_topological_order(course_id)
        assert len(order) == len(graph.nodes)

        # 7. Verify Media Store Figure Retention (Member 3 Contract)
        figures = media_store.list_figures(course_id)
        assert len(figures) >= 1
        fig = figures[0]
        assert fig.source.type == "pdf"
        assert fig.source.page == 1
        assert media_store.get_image_path(fig.image_rel_path).exists()
